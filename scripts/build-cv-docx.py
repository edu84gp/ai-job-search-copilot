import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE as RT

NAVY = RGBColor(0x1F, 0x38, 0x64)
GREY = RGBColor(0x40, 0x40, 0x40)
LINK = RGBColor(0x05, 0x63, 0xC1)
FONT = "Calibri"
INLINE = re.compile(r"\*\*(.+?)\*\*|\[([^\]]+)\]\(([^)]+)\)")


def style_run(run, size, bold=False, color=None):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), FONT)


def add_hyperlink(paragraph, text, url, size):
    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rfonts = OxmlElement("w:rFonts")
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), FONT)
    rpr.append(rfonts)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    rpr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.append(underline)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size * 2)))
    rpr.append(sz)
    run.append(rpr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_inline(paragraph, text, size, bold=False, color=None):
    position = 0
    for match in INLINE.finditer(text):
        if match.start() > position:
            run = paragraph.add_run(text[position:match.start()])
            style_run(run, size, bold, color)
        if match.group(1) is not None:
            run = paragraph.add_run(match.group(1))
            style_run(run, size, True, color)
        else:
            add_hyperlink(paragraph, match.group(2), match.group(3), size)
        position = match.end()
    if position < len(text):
        run = paragraph.add_run(text[position:])
        style_run(run, size, bold, color)


def spacing(paragraph, before=0, after=1.5):
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.0


def build(source, target):
    lines = [line.rstrip("\n") for line in Path(source).read_text(encoding="utf-8").splitlines()]
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(1.5)
    normal.paragraph_format.line_spacing = 1.0
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), FONT)

    doc.core_properties.author = "Eduardo García Prieto"
    doc.core_properties.last_modified_by = "Eduardo García Prieto"
    doc.core_properties.title = "Eduardo García Prieto — Currículum"
    doc.core_properties.subject = "Currículum Junior Full-Stack Developer"
    doc.core_properties.keywords = "React, Node.js, PostgreSQL, APIs REST, Full-Stack"

    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(0.9)
    section.bottom_margin = Cm(0.9)
    section.left_margin = Cm(1.5)
    section.right_margin = Cm(1.5)

    header_index = 0
    for line in lines:
        stripped = line.strip()
        if not stripped:
            continue
        if header_index == 0 and stripped.startswith("# "):
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            spacing(paragraph, 0, 1)
            run = paragraph.add_run(stripped[2:])
            style_run(run, 20, True, NAVY)
            header_index = 1
            continue
        if header_index == 1:
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            spacing(paragraph, 1, 1)
            add_inline(paragraph, stripped, 12, True, GREY)
            header_index = 2
            continue
        if header_index in (2, 3):
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            spacing(paragraph, 1 if header_index == 2 else 0, 0)
            add_inline(paragraph, stripped, 10, False, GREY)
            header_index += 1
            continue

        if stripped.startswith("## "):
            paragraph = doc.add_paragraph()
            spacing(paragraph, 5, 1.5)
            run = paragraph.add_run(stripped[3:].upper())
            style_run(run, 12, True, NAVY)
        elif stripped.startswith("- "):
            paragraph = doc.add_paragraph()
            spacing(paragraph, 0, 0.5)
            add_inline(paragraph, "\u2022 " + stripped[2:], 11)
        elif stripped.startswith("**") and "**" in stripped[2:]:
            paragraph = doc.add_paragraph()
            spacing(paragraph, 2.5, 1.5)
            add_inline(paragraph, stripped, 11)
        else:
            paragraph = doc.add_paragraph()
            spacing(paragraph, 0, 1.5)
            add_inline(paragraph, stripped, 11)

    doc.save(target)


if __name__ == "__main__":
    build(sys.argv[1], sys.argv[2])
