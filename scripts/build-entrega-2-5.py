from pathlib import Path
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from reportlab.lib import colors
from reportlab.lib.enums import TA_JUSTIFY
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

OUT = Path(__file__).resolve().parent.parent / "docs" / "ejercicios" / "entregables"

NAVY = "1F3864"
LIGHT = "EDF1F8"
GREY = "595959"
FONT_DIR = Path("C:/Windows/Fonts")

TITULO = "Eduardo García Prieto"
CURSO = "Employment Workshop 2026 · Unidad 2 · Lección 2.5 Actividades prácticas"
FECHA = "29 de octubre de 2026"
DIRECCION = (
    "Busco un puesto de desarrollador full-stack junior. Vengo de trabajar en sectores muy "
    "distintos y he aprendido a programar construyendo cosas. Los sectores que ya conozco, "
    "logística, hostelería, producción y gestión ecuestre, son un punto de partida y no un "
    "límite: me ayudan a entender de qué habla un cliente, no a encerrarme en un nicho."
)

DAFO = {
    "Fortalezas": [
        (
            "Programo interfaces web y me preocupo por que las use cualquiera",
            "He construido tableros que se manejan solo con el teclado y he corregido fallos "
            "que yo mismo había metido días antes.",
        ),
        (
            "Sé montar sistemas de IA por piezas",
            "He montado un asistente que reparte el trabajo entre varios agentes conectados a "
            "servicios externos, cada uno con permisos limitados para que no pueda hacer más "
            "de lo que debe.",
        ),
        (
            "Documento lo que construyo",
            "Cuando termino un proyecto lo dejo explicado: qué hace, cómo se pone en marcha y "
            "qué límites tiene.",
        ),
    ],
    "Debilidades": [
        (
            "Todavía no he demostrado que sé programar sin ayuda de IA",
            "Casi todo lo que tengo publicado lo construí apoyándome en asistentes. Me falta "
            "una pieza hecha a mano, de principio a fin.",
        ),
        (
            "No tengo pruebas automáticas en ningún proyecto",
            "Nunca he escrito una batería de tests. Sé lo que son y para qué sirven, pero no "
            "puedo enseñar ninguno.",
        ),
        (
            "Mi inglés es B1",
            "Leo documentación sin problema, pero hablar en una entrevista técnica me cuesta.",
        ),
    ],
    "Oportunidades": [
        (
            "Hay ofertas junior que no piden experiencia",
            "De las cinco que analicé, dos no pedían ningún año trabajado.",
        ),
        (
            "Muchos negocios pequeños siguen gestionándose en papel",
            "He visto de primera mano cómo trabajan una hípica, un hotel, una fábrica y una "
            "empresa de transporte.",
        ),
        (
            "La IA aplicada al trabajo diario todavía es rara en las ofertas",
            "Ninguna de las cinco ofertas que leí hablaba de agentes ni de modelos de lenguaje.",
        ),
    ],
    "Amenazas": [
        (
            "Muchas ofertas piden uno o dos años de experiencia",
            "Tres de las cinco lo pedían, y es el requisito que más se repite.",
        ),
        (
            "Algunas exigen inglés B2",
            "Una lo exigía y otra lo valoraba. Cuando aparece, es eliminatorio.",
        ),
        (
            "Las pruebas técnicas se hacen sin IA y con el reloj en contra",
            "Me evalúan escribiendo código a mano, sin las herramientas con las que trabajo "
            "a diario.",
        ),
    ],
}

CAME = [
    (
        "Corregir",
        "Añadir pruebas automáticas a mi proyecto de Pokédex y dejarlo explicado en el repositorio",
        "2026-11-11",
        "Convierte una debilidad que solo yo conozco en algo que se puede ver.",
    ),
    (
        "Afrontar",
        "Comprobar cuatro cosas antes de aplicar a una oferta que pida dos años: que cubro dos "
        "tercios de los requisitos, que la experiencia es lo único que me falta, que no hay "
        "requisito eliminatorio y que puedo demostrarlo con proyectos",
        "2026-11-11",
        "Convierte un filtro que me cierra puertas en un criterio que decido yo.",
    ),
    (
        "Mantener",
        "Hablar con dos contactos de hípica, hostelería o logística para confirmar que el "
        "problema que quiero resolver existe de verdad",
        "2026-11-30",
        "Confirma o descarta mi idea con dos conversaciones y sin gastar nada.",
    ),
    (
        "Explotar",
        "Aplicar cada semana a todas las ofertas junior sin experiencia que cumplan ese criterio",
        "Todas las semanas",
        "Convierte una oportunidad del mercado en candidaturas reales.",
    ),
]

PIE = (
    "Análisis personal presentado como actividad práctica de la unidad 2 del "
    "Employment Workshop 2026."
)


def shade(cell, hex_color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def docx_run(paragraph, text, size, bold=False, color=None):
    run = paragraph.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    return run


def docx_base(doc):
    section = doc.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)


def docx_header(doc, titulo, subtitulo):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    docx_run(p, titulo, 20, True, NAVY)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    docx_run(p, subtitulo, 11, False, GREY)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(14)
    docx_run(p, f"{CURSO} · {FECHA}", 9, False, GREY)
    borde = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), NAVY)
    borde.append(bottom)
    p._p.get_or_add_pPr().append(borde)


def docx_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(6)
    docx_run(p, text, 13, True, NAVY)
    return p


def docx_items(cell, items):
    cell.paragraphs[0]._p.getparent().remove(cell.paragraphs[0]._p)
    for index, (titulo, evidencia) in enumerate(items, start=1):
        p = cell.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        docx_run(p, f"{index}. {titulo}", 10, True)
        q = cell.add_paragraph()
        q.paragraph_format.space_after = Pt(8)
        q.paragraph_format.left_indent = Cm(0.4)
        docx_run(q, evidencia, 8.5, False, GREY)


def docx_footer(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    docx_run(p, PIE, 8, False, GREY)


def build_dafo_docx(path):
    doc = Document()
    docx_base(doc)
    docx_header(doc, "DAFO personal", TITULO)
    docx_heading(doc, "Punto de partida")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(12)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    docx_run(p, DIRECCION, 10)

    docx_heading(doc, "Análisis interno")
    table = doc.add_table(rows=2, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for col, name in enumerate(("Fortalezas", "Debilidades")):
        cell = table.cell(0, col)
        shade(cell, NAVY)
        cell.paragraphs[0].paragraph_format.space_after = Pt(2)
        docx_run(cell.paragraphs[0], name.upper(), 10, True, "FFFFFF")
        docx_items(table.cell(1, col), DAFO[name])
    for row in table.rows:
        for cell in row.cells:
            cell.width = Cm(8.5)

    docx_heading(doc, "Análisis externo")
    table = doc.add_table(rows=2, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for col, name in enumerate(("Oportunidades", "Amenazas")):
        cell = table.cell(0, col)
        shade(cell, NAVY)
        cell.paragraphs[0].paragraph_format.space_after = Pt(2)
        docx_run(cell.paragraphs[0], name.upper(), 10, True, "FFFFFF")
        docx_items(table.cell(1, col), DAFO[name])
    for row in table.rows:
        for cell in row.cells:
            cell.width = Cm(8.5)

    docx_footer(doc)
    doc.save(path)


def build_came_docx(path):
    doc = Document()
    docx_base(doc)
    docx_header(doc, "Plan CAME", TITULO)
    docx_heading(doc, "Una acción por movimiento")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    docx_run(
        p,
        "Cada letra del CAME se cubre con una acción y una fecha. No es un plan nuevo: es la "
        "misma lista de tareas que ya tengo, mirada desde cuatro ángulos distintos.",
        10,
    )

    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for col, name in enumerate(("Letra", "Acción", "Fecha", "Para qué sirve")):
        cell = table.rows[0].cells[col]
        shade(cell, NAVY)
        cell.paragraphs[0].paragraph_format.space_after = Pt(2)
        docx_run(cell.paragraphs[0], name.upper(), 9, True, "FFFFFF")

    for letra, accion, plazo, efecto in CAME:
        cells = table.add_row().cells
        cells[0].paragraphs[0].paragraph_format.space_after = Pt(2)
        docx_run(cells[0].paragraphs[0], letra, 10, True, NAVY)
        cells[1].paragraphs[0].paragraph_format.space_after = Pt(2)
        docx_run(cells[1].paragraphs[0], accion, 10)
        cells[2].paragraphs[0].paragraph_format.space_after = Pt(2)
        docx_run(cells[2].paragraphs[0], plazo, 9, False, GREY)
        cells[3].paragraphs[0].paragraph_format.space_after = Pt(2)
        docx_run(cells[3].paragraphs[0], efecto, 9, False, GREY)

    widths = (Cm(2.2), Cm(6.6), Cm(3.0), Cm(5.2))
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = widths[index]

    docx_footer(doc)
    doc.save(path)


def pdf_fonts():
    pdfmetrics.registerFont(TTFont("Calibri", str(FONT_DIR / "calibri.ttf")))
    pdfmetrics.registerFont(TTFont("Calibri-Bold", str(FONT_DIR / "calibrib.ttf")))
    pdfmetrics.registerFont(TTFont("Calibri-Italic", str(FONT_DIR / "calibrii.ttf")))
    pdfmetrics.registerFontFamily(
        "Calibri", normal="Calibri", bold="Calibri-Bold", italic="Calibri-Italic"
    )


def pdf_styles():
    return {
        "titulo": ParagraphStyle(
            "titulo", fontName="Calibri-Bold", fontSize=20, leading=24,
            textColor=colors.HexColor("#" + NAVY), spaceAfter=2,
        ),
        "subtitulo": ParagraphStyle(
            "subtitulo", fontName="Calibri", fontSize=11, leading=14,
            textColor=colors.HexColor("#" + GREY), spaceAfter=2,
        ),
        "meta": ParagraphStyle(
            "meta", fontName="Calibri", fontSize=9, leading=12,
            textColor=colors.HexColor("#" + GREY), spaceAfter=14,
        ),
        "seccion": ParagraphStyle(
            "seccion", fontName="Calibri-Bold", fontSize=13, leading=16,
            textColor=colors.HexColor("#" + NAVY), spaceBefore=10, spaceAfter=6,
        ),
        "cuerpo": ParagraphStyle(
            "cuerpo", fontName="Calibri", fontSize=10, leading=14,
            alignment=TA_JUSTIFY, spaceAfter=10,
        ),
        "item": ParagraphStyle(
            "item", fontName="Calibri-Bold", fontSize=10, leading=13, spaceAfter=1,
        ),
        "evidencia": ParagraphStyle(
            "evidencia", fontName="Calibri", fontSize=8.5, leading=11,
            textColor=colors.HexColor("#" + GREY), leftIndent=10, spaceAfter=8,
        ),
        "celda": ParagraphStyle(
            "celda", fontName="Calibri", fontSize=9, leading=12,
            textColor=colors.HexColor("#" + GREY),
        ),
        "celda_fuerte": ParagraphStyle(
            "celda_fuerte", fontName="Calibri-Bold", fontSize=10, leading=13,
            textColor=colors.HexColor("#" + NAVY),
        ),
        "pie": ParagraphStyle(
            "pie", fontName="Calibri", fontSize=8, leading=11,
            textColor=colors.HexColor("#" + GREY), spaceBefore=16,
        ),
    }


def pdf_header(styles, titulo):
    return [
        Paragraph(titulo, styles["titulo"]),
        Paragraph(TITULO, styles["subtitulo"]),
        Paragraph(f"{CURSO} · {FECHA}", styles["meta"]),
        Table(
            [[""]],
            colWidths=[17 * cm],
            rowHeights=[0.06 * cm],
            style=TableStyle([("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#" + NAVY))]),
        ),
        Spacer(1, 10),
    ]


def pdf_items(items, styles):
    flow = []
    for index, (titulo, evidencia) in enumerate(items, start=1):
        flow.append(Paragraph(f"{index}. {titulo}", styles["item"]))
        flow.append(Paragraph(evidencia, styles["evidencia"]))
    return flow


def pdf_cuadro(nombres, styles):
    cabecera = [
        Paragraph(n.upper(), ParagraphStyle(
            "cab", fontName="Calibri-Bold", fontSize=10, leading=13,
            textColor=colors.white,
        ))
        for n in nombres
    ]
    cuerpo = [pdf_items(DAFO[n], styles) for n in nombres]
    table = Table(
        [cabecera, cuerpo],
        colWidths=[8.5 * cm, 8.5 * cm],
        style=TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + NAVY)),
            ("BACKGROUND", (0, 1), (-1, 1), colors.HexColor("#" + LIGHT)),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#B4C0D6")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 8),
            ("RIGHTPADDING", (0, 0), (-1, -1), 8),
            ("TOPPADDING", (0, 0), (-1, -1), 6),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ]),
    )
    return table


def build_dafo_pdf(path, styles):
    doc = SimpleDocTemplate(
        str(path), pagesize=A4,
        leftMargin=2 * cm, rightMargin=2 * cm,
        topMargin=2 * cm, bottomMargin=2 * cm,
        title="DAFO personal", author=TITULO,
    )
    flow = pdf_header(styles, "DAFO personal")
    flow.append(Paragraph("Punto de partida", styles["seccion"]))
    flow.append(Paragraph(DIRECCION, styles["cuerpo"]))
    flow.append(Paragraph("Análisis interno", styles["seccion"]))
    flow.append(pdf_cuadro(["Fortalezas", "Debilidades"], styles))
    flow.append(Spacer(1, 14))
    flow.append(Paragraph("Análisis externo", styles["seccion"]))
    flow.append(pdf_cuadro(["Oportunidades", "Amenazas"], styles))
    flow.append(Paragraph(PIE, styles["pie"]))
    doc.build(flow)


def build_came_pdf(path, styles):
    doc = SimpleDocTemplate(
        str(path), pagesize=A4,
        leftMargin=2 * cm, rightMargin=2 * cm,
        topMargin=2 * cm, bottomMargin=2 * cm,
        title="Plan CAME", author=TITULO,
    )
    flow = pdf_header(styles, "Plan CAME")
    flow.append(Paragraph("Una acción por movimiento", styles["seccion"]))
    flow.append(Paragraph(
        "Cada letra del CAME se cubre con una acción y una fecha. No es un plan nuevo: es la "
        "misma lista de tareas que ya tengo, mirada desde cuatro ángulos distintos.",
        styles["cuerpo"],
    ))

    cabecera = [
        Paragraph(t, ParagraphStyle(
            "cab", fontName="Calibri-Bold", fontSize=9, leading=12,
            textColor=colors.white,
        ))
        for t in ("LETRA", "ACCIÓN", "FECHA", "PARA QUÉ SIRVE")
    ]
    filas = [cabecera]
    for letra, accion, plazo, efecto in CAME:
        filas.append([
            Paragraph(letra, styles["celda_fuerte"]),
            Paragraph(accion, ParagraphStyle("a", fontName="Calibri", fontSize=9.5, leading=12)),
            Paragraph(plazo, styles["celda"]),
            Paragraph(efecto, styles["celda"]),
        ])

    table = Table(
        filas,
        colWidths=[2.2 * cm, 6.4 * cm, 3.0 * cm, 5.4 * cm],
        style=TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + NAVY)),
            ("BACKGROUND", (0, 1), (-1, -1), colors.HexColor("#" + LIGHT)),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#B4C0D6")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 7),
            ("RIGHTPADDING", (0, 0), (-1, -1), 7),
            ("TOPPADDING", (0, 0), (-1, -1), 7),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
        ]),
    )
    flow.append(table)
    flow.append(Paragraph(PIE, styles["pie"]))
    doc.build(flow)


def main():
    OUT.mkdir(parents=True, exist_ok=True)

    build_dafo_docx(OUT / "2-5-dafo-personal.docx")
    build_came_docx(OUT / "2-5-plan-came.docx")

    if "pdf" in sys.argv[1:]:
        pdf_fonts()
        styles = pdf_styles()
        build_dafo_pdf(OUT / "2-5-dafo-personal.pdf", styles)
        build_came_pdf(OUT / "2-5-plan-came.pdf", styles)

    for file in sorted(OUT.iterdir()):
        print(f"{file.name}: {file.stat().st_size} bytes")


if __name__ == "__main__":
    main()
